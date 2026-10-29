"""Shared python-docx helpers and project constants for all Word documents."""
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from common import FIGURES_DIR, TABLES_DIR

TITLE = "Air Quality Across India: A Big Data Analysis of CPCB Monitoring Data (2010-2023) using Hadoop and PySpark"
COURSE = "Big Data Analytics - Course Project"
# Team members are kept out of version control: team.json holds
# [{"name": ..., "prn": ..., "contribution": ...}, ...]. Placeholders are used without it.
_team_file = Path(__file__).with_name("team.json")
if _team_file.exists():
    TEAM = [(m["name"], m["prn"], m["contribution"]) for m in json.loads(_team_file.read_text())]
else:
    TEAM = [("Team member 1", "", "50%"), ("Team member 2", "", "50%")]
ACCENT = RGBColor(0x1C, 0x5C, 0xAB)


def t(name):
    return pd.read_csv(TABLES_DIR / f"{name}.csv")


# ------------------------------------------------------------------ docx helpers
def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def table(doc, df, widths=None, font=9):
    tbl = doc.add_table(rows=1, cols=len(df.columns))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, col in enumerate(df.columns):
        cell = tbl.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(col))
        run.bold = True; run.font.size = Pt(font); run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, "1C5CAB")
    for _, row in df.iterrows():
        cells = tbl.add_row().cells
        for i, v in enumerate(row):
            if isinstance(v, float):
                v = f"{v:,.2f}".rstrip("0").rstrip(".") if abs(v) < 1000 else f"{v:,.0f}"
            elif isinstance(v, int):
                v = f"{v:,}"
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(font)
    if widths:
        for row in tbl.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return tbl


def caption(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


FIG_NO = [0]


def figure(doc, name, text, width=15.5):
    FIG_NO[0] += 1
    doc.add_picture(str(FIGURES_DIR / f"{name}.png"), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, f"Figure {FIG_NO[0]}: {text}")


TAB_NO = [0]


def table_title(doc, text):
    TAB_NO[0] += 1
    p = doc.add_paragraph()
    r = p.add_run(f"Table {TAB_NO[0]}: {text}"); r.bold = True; r.font.size = Pt(9)


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2EF")
    pPr.append(shd)
    r = p.add_run(text.strip("\n"))
    r.font.name = "Consolas"; r.font.size = Pt(8)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def h(doc, text, level=1):
    para = doc.add_heading(text, level=level)
    for r in para.runs:
        r.font.color.rgb = ACCENT
    return para


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(6)
    FIG_NO[0] = TAB_NO[0] = 0
    return doc


def title_block(doc, kicker, title, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(kicker.upper())
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(subtitle).font.size = Pt(11)


def save(doc, path, title):
    cp = doc.core_properties
    cp.author = cp.last_modified_by = " & ".join(n for n, _, _ in TEAM)
    cp.title, cp.subject, cp.comments = title, COURSE, ""
    doc.save(path)
    print(f"[docx] {path.relative_to(path.parents[1])}")

